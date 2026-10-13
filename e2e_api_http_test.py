"""
End-to-end API test using real HTTP endpoints:
1. POST /api/group with `test_ramassage.xlsx` as multipart upload
2. POST /api/export with returned `grouped_data` JSON
3. Save the returned XLSX and perform basic checks (sheets, merged taxi cells, red fill)

Requires the API server to be running at http://127.0.0.1:8000
"""
import json
import sys
from pathlib import Path

import requests
from openpyxl import load_workbook

API_BASE = "http://127.0.0.1:8000"


def post_group(file_path: str):
    url = f"{API_BASE}/api/group"
    with open(file_path, "rb") as f:
        files = {"file": (Path(file_path).name, f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        resp = requests.post(url, files=files)
    resp.raise_for_status()
    return resp.json()


def post_export(grouped_data, out_path: str):
    url = f"{API_BASE}/api/export"
    headers = {"Content-Type": "application/json"}
    # The API expects the JSON body to be the array itself (not wrapped as {"grouped_data": ...})
    resp = requests.post(url, json=grouped_data, headers=headers)
    resp.raise_for_status()
    with open(out_path, "wb") as f:
        f.write(resp.content)


def inspect_workbook(path: str):
    wb = load_workbook(path)
    sheets = wb.sheetnames
    has_grouped = "Grouped Agents" in sheets
    has_summary = "Summary" in sheets

    ws = wb["Grouped Agents"] if has_grouped else wb[wb.sheetnames[0]]
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    try:
        taxi_col_idx = headers.index("Groupe Taxi") + 1
    except ValueError:
        taxi_col_idx = None

    merged_present = False
    if taxi_col_idx:
        for m in ws.merged_cells.ranges:
            if m.min_col == taxi_col_idx and m.max_col == taxi_col_idx and m.max_row > m.min_row:
                merged_present = True
                break

    red_found = False
    if taxi_col_idx:
        for row in range(2, min(ws.max_row + 1, 200)):
            cell = ws.cell(row, taxi_col_idx)
            rgb = None
            try:
                if cell.fill and cell.fill.fgColor:
                    rgb = cell.fill.fgColor.rgb
            except Exception:
                rgb = None
            if rgb and (rgb.endswith("C00000") or rgb.endswith("FF0000")):
                red_found = True
                break

    return {
        "sheets": sheets,
        "has_grouped_sheet": has_grouped,
        "has_summary_sheet": has_summary,
        "taxi_col_idx": taxi_col_idx,
        "taxi_merged": merged_present,
        "taxi_red_fill_found": red_found,
    }


if __name__ == "__main__":
    input_file = Path("test_ramassage.xlsx")
    if not input_file.exists():
        print("Missing test_ramassage.xlsx")
        sys.exit(2)

    print("Posting to /api/group...")
    grouped_resp = post_group(str(input_file))
    grouped = grouped_resp.get("grouped_data")
    if not grouped:
        print("/api/group returned no grouped_data")
        print(grouped_resp)
        sys.exit(1)

    out = Path("e2e_grouped_output_api.xlsx")
    if out.exists():
        out.unlink()
    print("Posting grouped_data to /api/export and saving file...")
    post_export(grouped, str(out))

    report = inspect_workbook(str(out))
    print("E2E HTTP report:\n", json.dumps(report, indent=2, ensure_ascii=False))
    if report["has_grouped_sheet"] and report["has_summary_sheet"] and report["taxi_merged"]:
        print("E2E HTTP EXPORT OK")
        sys.exit(0)
    else:
        print("E2E HTTP EXPORT FAILED")
        sys.exit(1)
