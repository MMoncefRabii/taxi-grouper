"""
End-to-end test: run grouping pipeline on `test_ramassage.xlsx`, export to Excel,
and validate resulting workbook formatting (sheet names, header color, merged taxi cells, summary sheet).
"""
from pathlib import Path
import tempfile
import sys

import pandas as pd
from openpyxl import load_workbook

from src.core.preprocessing import load_dispatch_excel, detect_columns, filter_dataframe
from src.core.grouper import group_agents
from src.adapters.exporter import export_excel
from src.config import TAXI_COLUMN


def run_e2e(input_path: str, output_path: str) -> dict:
    result = {"input": input_path, "output": output_path, "checks": {}}
    df = load_dispatch_excel(input_path)
    cols = detect_columns(df)
    df_filtered = filter_dataframe(df, cols, time_value=None, type_course=None, operation=None)

    grouped, warnings = group_agents(df_filtered, columns=cols, use_ml=True)
    export_excel(grouped, output_path, taxi_column=TAXI_COLUMN)

    # Inspect workbook
    wb = load_workbook(output_path)
    sheets = wb.sheetnames
    result["checks"]["sheets"] = sheets
    has_grouped = "Grouped Agents" in sheets
    has_summary = "Summary" in sheets
    result["checks"]["has_grouped_sheet"] = has_grouped
    result["checks"]["has_summary_sheet"] = has_summary

    ws = wb["Grouped Agents"] if has_grouped else wb[wb.sheetnames[0]]
    header_fill = ws.cell(1, 1).fill.fgColor.rgb if ws.cell(1, 1).fill and ws.cell(1, 1).fill.fgColor else None
    result["checks"]["header_fill"] = header_fill

    # Find taxi column index
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    try:
        taxi_col_idx = headers.index(TAXI_COLUMN) + 1
    except ValueError:
        taxi_col_idx = None
    result["checks"]["taxi_col_idx"] = taxi_col_idx

    # Check merged cells for taxi column
    merged_present = False
    if taxi_col_idx:
        for m in ws.merged_cells.ranges:
            if m.min_col == taxi_col_idx and m.max_col == taxi_col_idx and m.max_row > m.min_row:
                merged_present = True
                break
    result["checks"]["taxi_merged"] = merged_present

    # Basic formatting checks: at least one taxi cell with RED fill (exporter sets taxi cell to red fill)
    red_found = False
    if taxi_col_idx:
        for row in range(2, min(ws.max_row + 1, 200)):
            cell = ws.cell(row, taxi_col_idx)
            try:
                rgb = cell.fill.fgColor.rgb
            except Exception:
                rgb = None
            if rgb and rgb.endswith("C00000") or (rgb and rgb.endswith("FF0000")):
                red_found = True
                break
    result["checks"]["taxi_red_fill_found"] = red_found

    return result


if __name__ == "__main__":
    input_file = Path("test_ramassage.xlsx")
    if not input_file.exists():
        print("Missing test_ramassage.xlsx in workspace root")
        sys.exit(2)
    out = Path("e2e_grouped_output.xlsx")
    if out.exists():
        out.unlink()
    report = run_e2e(str(input_file), str(out))
    print("E2E report:\n", report)
    if report["checks"]["has_grouped_sheet"] and report["checks"]["has_summary_sheet"] and report["checks"]["taxi_merged"]:
        print("E2E EXPORT OK")
        sys.exit(0)
    else:
        print("E2E EXPORT FAILED: checks did not pass")
        sys.exit(1)
