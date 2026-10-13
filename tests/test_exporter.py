import pandas as pd
from openpyxl import load_workbook

from src.adapters.exporter import export_excel, export_like_historique


def test_exporter_creates_summary(tmp_path):
    df = pd.DataFrame(
        {
            "Nom": ["A", "B"],
            "Adresse": ["Ariana", "Ariana"],
            "Zone": ["H", "H"],
            "Zone Name": ["Ariana", "Ariana"],
            "Confidence": [1.0, 1.0],
            "Matched Alias": ["Ariana", "Ariana"],
            "Normalized Address": ["ariana", "ariana"],
            "Groupe Taxi": [1, 1],
        }
    )
    output = export_excel(df, tmp_path / "out.xlsx")
    wb = load_workbook(output)
    assert "Grouped Agents" in wb.sheetnames
    assert "Summary" in wb.sheetnames


def test_historique_export_uses_numero_prix_header(tmp_path):
    df = pd.DataFrame(
        {
            "Date": ["2026-05-29"],
            "Matricule": [1],
            "Nom - Prénom": ["Agent A"],
            "Opération": ["OP"],
            "Site": ["Site 1"],
            "Numero": ["555"],
            "Adresse": ["Oued Lil Manouba"],
            "Heure d'arrivée": ["07:00"],
            "PRIX": [12],
            "Zone": ["G"],
            "DetectedDestination": ["Oued Lil"],
            "GroupReason": ["Same destination"],
            "Groupe Taxi": [1],
        }
    )
    output = export_like_historique(df, str(tmp_path / "dispatch.xlsx"))
    wb = load_workbook(output)
    ws = wb.active
    headers = [ws.cell(1, col).value for col in range(1, 11)]
    assert headers == ["Date", "Matricule", "Nom - Prénom", "Opération", "Site", "Numéro", "Adresse", "Heure d'arrivée", "Taxi", "PRIX"]
    assert ws.cell(2, 6).value == "555"
    assert ws.cell(2, 10).value == 12
